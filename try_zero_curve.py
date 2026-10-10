"""
Turkish lira (TRY) government zero-coupon yield curve from bond-level data.

Inputs (Excel or CSV, in the formats you already have):
  definitions   ISIN | Getiri Türü (Sabit/Fixed, İskontolu/Discounted) | Yıllık Kupon Sıklığı
                | İtfa Tarihi (maturity) | Sonraki Kupon Oranı % (coupon PER PERIOD, e.g. 16.30 = 32.6%/yr)
  market data   [Trade Date] | ISIN | CLOSING (clean price per 100) | TOTAL TRADED VALUE
                (columns found by header; several rows per ISIN are combined into a value-weighted
                average price; the trade date is taken from the file when it has one date)

Method
  1. Clean prices: values exported as whole numbers with a thousands format (89805) are
     read as 89.805; Turkish decimal commas are handled.
  2. Universe: TL fixed-coupon Treasury bonds and discounted bills only. Excluded and listed:
     floaters (…F…), CPI-linked or other bonds missing from the definitions, all TRD… codes
     (gold/FX-linked Treasury issues and lease certificates, even when listed as "fixed"),
     corporate paper (TRF…), prices not quoted per 100 TL, and yields outside 0–200%.
  3. Cash flows: Turkish Treasury coupons fall every 182 days, counted back from maturity.
     Accrued interest = coupon × days since last coupon / days in period. Dirty = clean + accrued.
  4. Market yields in the Turkish quoting convention (checked against Gedik's DİBS bulletin):
     periodic rate over whole coupon periods, compounded to (1+r)^2-1; Act/365 for bills and bonds
     in their final period. Plus simple yield and Macaulay duration.
  5. Zero curve: Nelson–Siegel (default) or Svensson fitted to dirty prices discounted on actual
     days, errors weighted by 1/duration (≈ yield errors), multi-start least squares. Output as annual-compound and continuous
     zero rates, par yields and 1-year forward rates.
  Optional: --us adds the US Treasury zero curve for the same day (via ust_zero_curve.py).

Usage
  python try_zero_curve.py --defs definitions.xlsx --market Market_Data.xlsx       # date from the file
  python try_zero_curve.py --defs definitions.xlsx --market Market_Data.xlsx --date 2026-10-08
  python try_zero_curve.py ... --settle-days 1          # T+1 settlement instead of same day
  python try_zero_curve.py ... --us                     # add the US curve panel (live Treasury data)
  python try_zero_curve.py ... --us --us-par-csv par.csv  # US curve from a saved Treasury CSV

Outputs (folder try_curve_<date>/): try_zero_curve.png, try_zero_curve.csv and try_bonds.csv
(both start with a date column), try_excluded.csv, try_fit.json
"""

import argparse
import json
import re
from datetime import date, timedelta
from pathlib import Path

import numpy as np
import pandas as pd
from scipy.optimize import least_squares

PERIOD_DAYS = 182                     # Turkish Treasury coupon period
STD = {"3M": .25, "6M": .5, "1Y": 1, "2Y": 2, "3Y": 3, "4Y": 4, "5Y": 5, "7Y": 7, "10Y": 10}


# ============================================================ reading inputs
def _read_any(path):
    p = Path(path)
    if p.suffix.lower() in (".xlsx", ".xlsm", ".xls"):
        return p
    return pd.read_csv(p, dtype=str)


def parse_number(v, fmt=None):
    """Excel cell -> float. Whole numbers shown with a thousands format are prices that lost their
    decimal point on export (89805 -> 89.805). Accepts '89,87' and '1.234,5' styles."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return np.nan
    if isinstance(v, (int, np.integer)) and fmt and "#,##0" in fmt and "." not in fmt.split(";")[0]:
        return v / 1000.0
    if isinstance(v, (int, float, np.integer, np.floating)):
        return float(v)
    s = str(v).strip().replace(" ", "")
    if "," in s and "." in s:                    # 1.234,56 (TR) or 1,234.56 (US)
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return np.nan


def _parse_date(v):
    """Excel date cell or text ('08.10.2026' day-first, '2026-10-08') -> date."""
    if v is None or (isinstance(v, float) and np.isnan(v)) or (isinstance(v, str) and not v.strip()):
        return None
    if hasattr(v, "date") and not isinstance(v, str):
        return v.date() if hasattr(v, "hour") else v
    s = str(v).strip()
    ts = pd.to_datetime(s, dayfirst=not s[:4].isdigit(), errors="coerce")
    return None if pd.isna(ts) else ts.date()


def _market_columns(header):
    """Find column positions by header name, so column order doesn't matter."""
    pos = {}
    for i, h in enumerate(header):
        lh = str(h or "").strip().lower()
        if "isin" in lh:
            pos.setdefault("isin", i)
        elif "date" in lh or "tarih" in lh:
            pos.setdefault("date", i)
        elif "clos" in lh or "kapan" in lh or "price" in lh or "fiyat" in lh:
            pos.setdefault("price", i)
        elif "value" in lh or "hacim" in lh or "tutar" in lh:
            pos.setdefault("value", i)
    if "isin" not in pos or "price" not in pos:            # no recognisable header: old fixed layout
        pos = {"isin": 0, "price": 1, "value": 2}
    return pos


def read_market(path, trade_date=None):
    """Rows of (trade date, ISIN, clean price, traded value) -> one value-weighted price per ISIN.
    Columns are found by header ('Trade Date', 'ISIN', 'CLOSING', 'TOTAL TRADED VALUE'); the date
    column is optional. Returns (prices, trade_date). With several dates in the file, pass trade_date."""
    src = _read_any(path)
    rows = []
    if isinstance(src, Path):
        import openpyxl
        ws = openpyxl.load_workbook(src, data_only=True).active
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        pos = _market_columns(header)
        for r in ws.iter_rows(min_row=2):
            if r[pos["isin"]].value is None:
                continue
            pc = r[pos["price"]]
            rows.append((_parse_date(r[pos["date"]].value) if "date" in pos else None,
                         str(r[pos["isin"]].value).strip(), parse_number(pc.value, pc.number_format),
                         parse_number(r[pos["value"]].value) if "value" in pos and pos["value"] < len(r) else np.nan))
    else:
        pos = _market_columns(list(src.columns))
        for _, r in src.iterrows():
            rows.append((_parse_date(r.iloc[pos["date"]]) if "date" in pos else None,
                         str(r.iloc[pos["isin"]]).strip(), parse_number(r.iloc[pos["price"]]),
                         parse_number(r.iloc[pos["value"]]) if "value" in pos else np.nan))
    m = pd.DataFrame(rows, columns=["trade_date", "isin", "price", "value"]).dropna(subset=["price"])
    m["value"] = m["value"].fillna(0.0)

    dates = sorted({d for d in m["trade_date"] if d is not None})
    if trade_date is not None:
        if dates:
            if trade_date not in dates:
                raise SystemExit(f"No prices for {trade_date} in {Path(path).name}; dates in file: "
                                 + ", ".join(map(str, dates)))
            m = m[m["trade_date"] == trade_date]
    elif len(dates) == 1:
        trade_date = dates[0]
    elif len(dates) > 1:
        raise SystemExit(f"{Path(path).name} has several trade dates ({', '.join(map(str, dates))}); pick one with --date")
    else:
        raise SystemExit(f"{Path(path).name} has no trade-date column; pass --date YYYY-MM-DD")

    def agg(g):
        w = g["value"] if g["value"].sum() > 0 else None
        return pd.Series({"clean": float(np.average(g["price"], weights=w)), "trades": len(g),
                          "traded_value": g["value"].sum(), "px_min": g["price"].min(), "px_max": g["price"].max()})
    return m.groupby("isin").apply(agg).reset_index(), trade_date


def read_definitions(path):
    src = _read_any(path)
    d = pd.read_excel(src) if isinstance(src, Path) else src
    d = d.dropna(how="all")
    cols = {}
    for c in d.columns:
        lc = str(c).lower()
        if "isin" in lc:
            cols[c] = "isin"
        elif "getiri" in lc or "type" in lc:
            cols[c] = "kind"
        elif "sıklı" in lc or "frequency" in lc or "sikli" in lc:
            cols[c] = "freq"
        elif "itfa" in lc or "maturity" in lc:
            cols[c] = "maturity"
        elif "kupon oran" in lc or "coupon rate" in lc:
            cols[c] = "coupon_pp"
    d = d.rename(columns=cols)
    need = {"isin", "kind", "freq", "maturity"}
    if not need <= set(d.columns):
        raise SystemExit(f"definitions file is missing columns: {need - set(d.columns)}")
    d["isin"] = d["isin"].astype(str).str.strip().str.upper()
    d = d[d["isin"].str.match(r"^[A-Z]{2}[A-Z0-9]{9}[0-9]$")]          # drop blank/footer rows
    # values may be text (tbliste_import.py reads with dtype=str): parse each cell
    d["maturity"] = d["maturity"].map(_parse_maturity)
    d = d[d["maturity"].notna()]
    d["freq"] = d["freq"].map(lambda v: parse_number(v)).fillna(0).astype(int)
    d["coupon_pp"] = d["coupon_pp"].map(lambda v: parse_number(v)) if "coupon_pp" in d else np.nan
    d["kind"] = d["kind"].astype(str).str.strip()
    return d.drop_duplicates("isin")


def _parse_maturity(v):
    """Excel date, Excel serial number, '13.11.2030', '13/11/2030' or '2030-11-13 00:00:00' -> date."""
    if isinstance(v, (int, float, np.integer, np.floating)) and not pd.isna(v) and 20000 < float(v) < 80000:
        return (pd.Timestamp("1899-12-30") + pd.Timedelta(days=float(v))).date()
    if isinstance(v, str):
        s = v.strip()
        if re.fullmatch(r"\d{5}(\.0+)?", s) and 20000 < float(s) < 80000:
            return (pd.Timestamp("1899-12-30") + pd.Timedelta(days=float(s))).date()
        v = s.replace("/", ".") if re.fullmatch(r"\d{1,2}/\d{1,2}/\d{4}", s) else s
    return _parse_date(v)


# ============================================================ bond maths
def schedule(maturity: date, settle: date, freq: int):
    """Coupon dates after settlement and the last coupon date on/before it (182-day steps from maturity)."""
    step = timedelta(days=round(365 / freq)) if freq not in (2,) else timedelta(days=PERIOD_DAYS)
    dates = [maturity]
    while dates[-1] > settle:
        dates.append(dates[-1] - step)
    prev = dates[-1]
    future = sorted(x for x in dates if x > settle)
    return future, prev, step.days


def cashflows(row, settle):
    """(times in years Act/365, amounts per 100, accrued interest)."""
    if row["freq"] == 0:
        t = (row["maturity"] - settle).days / 365.0
        return np.array([t]), np.array([100.0]), 0.0
    c = float(row["coupon_pp"])
    future, prev, period = schedule(row["maturity"], settle, row["freq"])
    accrued = c * (settle - prev).days / period
    t = np.array([(x - settle).days / 365.0 for x in future])
    amt = np.full(len(t), c)
    amt[-1] += 100.0
    return t, amt, accrued


def ytm_compound(dirty, t, amt):
    """Annual-compound Act/365 yield: dirty = Σ CF / (1+y)^t."""
    f = lambda y: np.sum(amt / (1 + y) ** t) - dirty
    lo, hi = -0.5, 5.0
    for _ in range(200):
        mid = (lo + hi) / 2
        if f(mid) > 0:
            lo = mid
        else:
            hi = mid
    return (lo + hi) / 2


def ytm_quote(dirty, row, settle):
    """Yield as quoted in the Turkish market (matches Borsa Istanbul / broker bulletins):
    bills and bonds in their final period: annual compound on Act/365;
    otherwise a periodic rate r over whole coupon periods (first one fractional), quoted as (1+r)^freq - 1."""
    if row["freq"] == 0:
        t, amt, _ = cashflows(row, settle)
        return ytm_compound(dirty, t, amt)
    future, prev, period = schedule(row["maturity"], settle, row["freq"])
    c = float(row["coupon_pp"])
    if len(future) == 1:
        days = (future[0] - settle).days
        return ((c + 100.0) / dirty) ** (365.0 / days) - 1
    e = np.array([(future[0] - settle).days / period + k for k in range(len(future))])
    amt = np.full(len(future), c)
    amt[-1] += 100.0
    lo, hi = -0.5, 3.0
    for _ in range(200):
        mid = (lo + hi) / 2
        if np.sum(amt / (1 + mid) ** e) > dirty:
            lo = mid
        else:
            hi = mid
    return (1 + (lo + hi) / 2) ** row["freq"] - 1


def macaulay(y, t, amt):
    pv = amt / (1 + y) ** t
    return float(np.sum(t * pv) / np.sum(pv))


# ============================================================ Nelson-Siegel-Svensson
def nss_cc(p, t):
    """Continuously compounded zero rate (decimal) at maturity t (years)."""
    b0, b1, b2, b3, l1, l2 = p
    t = np.maximum(np.asarray(t, float), 1e-6)
    x1, x2 = t / l1, t / l2
    f1 = (1 - np.exp(-x1)) / x1
    return b0 + b1 * f1 + b2 * (f1 - np.exp(-x1)) + b3 * ((1 - np.exp(-x2)) / x2 - np.exp(-x2))


def fit_curve(bonds, model="ns", seed=0):
    """model 'ns' = Nelson-Siegel (4 params, default: stable with few bonds); 'nss' = Svensson (6 params)."""
    full = (lambda p: p) if model == "nss" else (lambda p: np.r_[p[:3], 0.0, p[3], 1.0])

    def resid(q):
        p = full(q)
        out = []
        for _, b in bonds.iterrows():
            model = np.sum(b["amt"] * np.exp(-nss_cc(p, b["t"]) * b["t"]))
            out.append((model - b["dirty"]) / (b["dirty"] * max(b["duration"], 0.05)))   # ≈ yield error
        return np.array(out)

    y = bonds["ytm_cc"].to_numpy()
    rng = np.random.default_rng(seed)
    best = None
    if model == "nss":
        lb, ub = [0.0, -1.0, -2.0, -2.0, 0.05, 0.5], [2.0, 2.0, 2.0, 2.0, 10.0, 30.0]
        starts = [[y[-1], y[0] - y[-1], 0.0, 0.0, 1.0, 5.0]]
        starts += [[y[-1] + rng.normal(0, .02), y[0] - y[-1] + rng.normal(0, .03), rng.normal(0, .1),
                    rng.normal(0, .1), rng.uniform(.2, 4), rng.uniform(3, 12)] for _ in range(60)]
    else:
        lb, ub = [0.0, -1.0, -2.0, 0.05], [2.0, 2.0, 2.0, 10.0]
        starts = [[y[-1], y[0] - y[-1], 0.0, 1.5]]
        starts += [[rng.uniform(0, .6), rng.uniform(-.5, .5), rng.uniform(-1, 1), rng.uniform(.1, 8)] for _ in range(60)]
    for s in starts:
        s = np.clip(s, np.array(lb) + 1e-6, np.array(ub) - 1e-6)
        try:
            r = least_squares(resid, s, bounds=(lb, ub), method="trf")
        except Exception:
            continue
        if best is None or r.cost < best.cost:
            best = r
    return full(best.x), resid(best.x)


# ============================================================ main pipeline
def build(defs_path, market_path, trade_date: date = None, settle_days=0, model="ns", log=print):
    """trade_date: optional when the market file has a single Trade Date; required to pick one of several."""
    d = read_definitions(defs_path)
    m, trade_date = read_market(market_path, trade_date)
    settle = trade_date + timedelta(days=settle_days)
    j = m.merge(d, on="isin", how="left")

    def reason(r):
        isin = r["isin"]
        if isin.startswith("TRD"):                     # checked first: some are listed as "fixed" Treasury issues
            return "TRD: gold/FX-linked or lease certificate, not a TL nominal bond"
        if isin.startswith("TRF") or (isin.startswith("TRS") and not isin.startswith("TRST")):
            return "corporate / bank paper, not a Treasury bond"
        if pd.isna(r.get("kind")):
            if isin[9:10] == "F":
                return "floating-rate note"
            return "not in definitions (CPI-linked or floating?)"
        if not (30.0 <= r["clean"] <= 200.0):
            return f"price {r['clean']:g} not quoted per 100 TL"
        if r["maturity"] <= settle:
            return "matured"
        if r["isin"][9:10] in ("K", "A") and r["freq"] != 0:
            return "STRIPS: add as discounted to use"
        if r["freq"] > 0 and pd.isna(r["coupon_pp"]):
            return "no coupon rate"
        if not (r["kind"].startswith("Sabit") or r["kind"].startswith("İsk") or "Fixed" in r["kind"] or "Discount" in r["kind"]):
            return f"type {r['kind']}"
        return ""

    j["excluded"] = j.apply(reason, axis=1)
    use = j[j["excluded"] == ""].copy()
    if len(use) < 6:
        raise SystemExit(f"Only {len(use)} usable bonds; need at least 6 for a Nelson–Siegel–Svensson fit.")

    recs = []
    for _, r in use.iterrows():
        t, amt, acc = cashflows(r, settle)
        dirty = r["clean"] + acc
        y = ytm_compound(dirty, t, amt)                 # actual-time yield (used for duration / fit seed)
        yq = ytm_quote(dirty, r, settle)                # market-convention compound yield
        T = t[-1]
        recs.append(dict(isin=r["isin"], kind="bill" if r["freq"] == 0 else "fixed", maturity=r["maturity"],
                         years=T, coupon_pp=r["coupon_pp"], coupon_annual=(r["coupon_pp"] * r["freq"]) if r["freq"] else 0.0,
                         clean=r["clean"], accrued=acc, dirty=dirty, trades=int(r["trades"]),
                         traded_value=r["traded_value"], ytm=yq, simple=((1 + y) ** T - 1) / T,
                         row=dict(freq=r["freq"], coupon_pp=r["coupon_pp"], maturity=r["maturity"]),
                         ytm_cc=np.log1p(y), duration=macaulay(y, t, amt), t=t, amt=amt))
    bonds = pd.DataFrame(recs).sort_values("years").reset_index(drop=True)
    # safety net: a yield outside 0–200% means the price or the definition is wrong for this bond
    bad = (bonds["ytm"] <= 0) | (bonds["ytm"] >= 2.0)
    if bad.any():
        for _, b in bonds[bad].iterrows():
            log(f"  excluded {b['isin']}: implausible yield {b['ytm'] * 100:.1f}% (clean {b['clean']:g})")
            j.loc[j["isin"] == b["isin"], "excluded"] = f"implausible yield {b['ytm'] * 100:.1f}%"
        bonds = bonds[~bad].reset_index(drop=True)
        if len(bonds) < 6:
            raise SystemExit(f"Only {len(bonds)} usable bonds after removing outliers; need at least 6.")
    log(f"Settlement {settle}: {len(bonds)} bonds used, {int((j['excluded'] != '').sum())} excluded")

    params, res = fit_curve(bonds, model)
    bonds["fit_err_bp"] = res * 1e4
    bonds["fitted_dirty"] = [float(np.sum(b["amt"] * np.exp(-nss_cc(params, b["t"]) * b["t"]))) for _, b in bonds.iterrows()]
    bonds["fitted_ytm"] = [ytm_quote(fd, b["row"], settle) for fd, (_, b) in zip(bonds["fitted_dirty"], bonds.iterrows())]
    bonds["ytm_err_bp"] = (bonds["ytm"] - bonds["fitted_ytm"]) * 1e4
    rmse = float(np.sqrt(np.mean(bonds["ytm_err_bp"] ** 2)))
    log(f"{'Nelson-Siegel-Svensson' if model == 'nss' else 'Nelson-Siegel'} fit: yield RMSE {rmse:.0f} bp, "
        f"max |error| {bonds['ytm_err_bp'].abs().max():.0f} bp")

    tmax = float(bonds["years"].max())
    grid = np.round(np.arange(0.25, max(10.0, np.ceil(tmax)) + 1e-9, 0.25), 4)
    zc = nss_cc(params, grid)
    curve = pd.DataFrame({"years": grid, "zero_cc": zc, "zero_annual": np.expm1(zc)})
    # par yield of a new semi-annual coupon bond, quoted like the market: (1 + c/2)^2 - 1.
    # Computed at whole half-years (no accrued interest), interpolated in between.
    halves = np.arange(0.5, grid.max() + 1e-9, 0.5)
    par_h = []
    for T in halves:
        ts = np.arange(0.5, T + 1e-9, 0.5)
        dfs = np.exp(-nss_cc(params, ts) * ts)
        c = 2 * (1 - dfs[-1]) / dfs.sum()
        par_h.append((1 + c / 2) ** 2 - 1)
    curve["par_compound"] = np.interp(grid, np.r_[0.0, halves], np.r_[np.expm1(nss_cc(params, 1e-4)), par_h])
    curve.loc[curve["years"] < 0.5, "par_compound"] = curve["zero_annual"]
    z1 = nss_cc(params, grid + 1.0)
    curve["fwd_1y_annual"] = np.expm1((z1 * (grid + 1) - zc * grid))
    curve["extrapolated"] = (grid < bonds["years"].min() - 1e-9) | (grid > tmax + 1e-9)

    excluded = j[j["excluded"] != ""][["isin", "clean", "excluded"]].sort_values("isin")
    return dict(settle=settle, trade_date=trade_date, bonds=bonds, curve=curve, params=params, model=model,
                rmse_bp=rmse, excluded=excluded, tmin=float(bonds["years"].min()), tmax=tmax)


def us_curve(day: date, par_csv=None, log=print):
    """US Treasury zero curve (annual compounding, %) for `day` or the latest day before it."""
    import ust_zero_curve as uzc
    if par_csv:
        df = pd.read_csv(par_csv)
        df["Date"] = pd.to_datetime(df["Date"], format="%m/%d/%Y")
        df = df.set_index("Date").sort_index()
        mats = {c: uzc.tenor_to_years(c) for c in df.columns}
        df = df[[c for c, t in mats.items() if t is not None]]
        df.columns = [mats[c] for c in df.columns]
        par = df[sorted(df.columns)]
    else:
        par = uzc.load_treasury({day.year} | ({day.year - 1} if day.month == 1 else set()))
    par = par[par.index <= pd.Timestamp(day)]
    if par.empty:
        raise ValueError("no US curve on or before that date")
    used = par.index[-1].date()
    if used != day:
        log(f"US curve: {day} not published yet; using {used}")
    z_semi = uzc.bootstrap_zero(par.iloc[-1])                   # % semi-annual
    z_annual = ((1 + z_semi / 200) ** 2 - 1) * 100
    return used, z_annual


# ============================================================ outputs
def chart(res, path, us=None):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    b, c = res["bonds"], res["curve"]
    ink2, rule, main, alt, warm = "#4a5160", "#e1e4ea", "#c0392b", "#2a78d6", "#d95926"
    ncol = 3 if us is not None else 2
    fig, axs = plt.subplots(1, ncol, figsize=(6.2 * ncol, 5.6),
                            gridspec_kw={"width_ratios": [1.5, 1, 1.1][:ncol]})
    ax = axs[0]
    inr = ~c["extrapolated"]
    ax.plot(c["years"][inr], c["zero_annual"][inr] * 100, color=main, lw=2.6, label="Zero rate (annual comp.)")
    ax.plot(c["years"], c["zero_annual"] * 100, color=main, lw=1.2, ls=":", alpha=.8)
    ax.plot(c["years"][inr], c["par_compound"][inr] * 100, color=ink2, lw=1.4, ls="--", label="Par yield (compound, as quoted)")
    ax.plot(c["years"][inr], c["fwd_1y_annual"][inr] * 100, color="#8e6bbf", lw=1.2, alpha=.8, label="1-year forward")
    for kind, mk in (("bill", "s"), ("fixed", "o")):
        s = b[b["kind"] == kind]
        ax.scatter(s["years"], s["ytm"] * 100, s=np.clip(np.sqrt(s["traded_value"]) / 900, 18, 160),
                   marker=mk, facecolor="white", edgecolor=main, lw=1.4, zorder=3,
                   label=f"Market compound yield ({'bills' if kind == 'bill' else 'fixed coupon'})")
    ax.set_xlabel("Maturity (years)")
    ax.set_ylabel("%")
    ax.set_title(f"Turkish lira government zero curve · settlement {res['settle']:%d %b %Y}", loc="left", fontsize=12)
    ax.grid(alpha=.35, color=rule)
    ax.legend(frameon=False, fontsize=8.5, loc="upper right")
    ax.text(0.01, 0.02, f"{'Nelson–Siegel–Svensson' if res['model'] == 'nss' else 'Nelson–Siegel'} fit to {len(b)} bonds · yield RMSE {res['rmse_bp']:.0f} bp · "
            f"dotted = outside {res['tmin']:.1f}–{res['tmax']:.1f}y of data",
            transform=ax.transAxes, fontsize=8, color=ink2)

    ax = axs[1]
    ax.barh(range(len(b)), b["ytm_err_bp"], color=[alt if v < 0 else warm for v in b["ytm_err_bp"]], height=.6)
    ax.set_yticks(range(len(b)), [f"{r.isin}  {r.maturity:%b %y}" for r in b.itertuples()], fontsize=8)
    ax.invert_yaxis()
    ax.axvline(0, color=ink2, lw=.8)
    ax.set_title("Market − fitted yield, bp", loc="left", fontsize=11)
    ax.grid(axis="x", alpha=.35, color=rule)

    if us is not None:
        used, zus = us
        ax = axs[2]
        t = np.array(c["years"][inr])
        ax.plot(t, c["zero_annual"][inr] * 100, color=main, lw=2.4, label=f"TRY ({res['settle']:%d %b})")
        zu = np.interp(t, zus.index, zus.values)
        ax.plot(t, zu, color=alt, lw=2.4, label=f"USD ({used:%d %b})")
        ax.fill_between(t, zu, c["zero_annual"][inr] * 100, color=main, alpha=.06)
        mid = t[len(t) // 2]
        ax.annotate(f"spread ≈ {np.interp(mid, t, c['zero_annual'][inr] * 100) - np.interp(mid, t, zu):.1f} pts at {mid:.0f}y",
                    (mid, (np.interp(mid, t, c['zero_annual'][inr] * 100) + np.interp(mid, t, zu)) / 2),
                    ha="center", fontsize=9, color=ink2)
        ax.set_title("TRY vs USD zero curves (annual comp.)", loc="left", fontsize=11)
        ax.set_xlabel("Maturity (years)")
        ax.set_ylabel("%")
        ax.grid(alpha=.35, color=rule)
        ax.legend(frameon=False, fontsize=9)
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)


def main():
    ap = argparse.ArgumentParser(description="TRY government zero curve from bond prices.")
    ap.add_argument("--defs", required=True, help="definitions file (xlsx/csv)")
    ap.add_argument("--market", required=True, help="market data file (xlsx/csv)")
    ap.add_argument("--date", help="trade date YYYY-MM-DD (or 08.10.2026); optional when the market file "
                                   "has a single Trade Date, required to pick one of several")
    ap.add_argument("--settle-days", type=int, default=0, help="settlement lag in days (0 = same day)")
    ap.add_argument("--model", choices=["ns", "nss"], default="ns",
                    help="ns = Nelson-Siegel (default, stable with ~15 bonds); nss = Svensson (more bonds needed)")
    ap.add_argument("--us", action="store_true", help="add the US Treasury zero curve panel")
    ap.add_argument("--us-par-csv", help="Treasury par-curve CSV instead of downloading")
    ap.add_argument("--out-dir")
    a = ap.parse_args()

    res = build(a.defs, a.market, _parse_date(a.date) if a.date else None, a.settle_days, a.model)
    out = Path(a.out_dir or f"try_curve_{res['trade_date']}")
    out.mkdir(parents=True, exist_ok=True)

    us = None
    if a.us:
        try:
            us = us_curve(res["settle"], a.us_par_csv)
        except Exception as ex:
            print(f"US curve skipped: {type(ex).__name__}: {ex}")
    chart(res, out / "try_zero_curve.png", us)

    c = res["curve"]
    std = pd.DataFrame({"date": res["trade_date"].isoformat(), "tenor": list(STD), "years": list(STD.values())})
    for col in ("zero_annual", "zero_cc", "par_compound", "fwd_1y_annual"):
        std[col + "_pct"] = np.interp(std["years"], c["years"], c[col]) * 100
    std["extrapolated"] = (std["years"] < res["tmin"]) | (std["years"] > res["tmax"])
    if us is not None:
        std["usd_zero_annual_pct"] = np.interp(std["years"], us[1].index, us[1].values)
        std["spread_pts"] = std["zero_annual_pct"] - std["usd_zero_annual_pct"]
    std.round(4).to_csv(out / "try_zero_curve.csv", index=False)
    cols = ["isin", "kind", "maturity", "years", "coupon_pp", "coupon_annual", "clean", "accrued", "dirty", "trades",
            "traded_value", "ytm", "simple", "duration", "fitted_ytm", "ytm_err_bp"]
    bt = res["bonds"][cols].copy()
    bt.insert(0, "date", res["trade_date"].isoformat())
    for k in ("ytm", "simple", "fitted_ytm"):
        bt[k] = bt[k] * 100
    bt.round(4).to_csv(out / "try_bonds.csv", index=False)
    res["excluded"].to_csv(out / "try_excluded.csv", index=False)
    (out / "try_fit.json").write_text(json.dumps({
        "trade_date": str(res["trade_date"]), "settlement": str(res["settle"]), "model": ("Nelson-Siegel-Svensson" if res["model"] == "nss" else "Nelson-Siegel") + ", continuous compounding",
        "params": dict(zip(["beta0", "beta1", "beta2", "beta3", "tau1", "tau2"], map(float, res["params"]))),
        "yield_rmse_bp": res["rmse_bp"], "bonds_used": int(len(res["bonds"])), "data_range_years": [res["tmin"], res["tmax"]],
        "us_curve_date": str(us[0]) if us is not None else None}, indent=1))

    print("\nStandard tenors (%):")
    print(std.round(2).to_string(index=False))
    print(f"\nWritten to {out.resolve()}")


if __name__ == "__main__":
    main()
